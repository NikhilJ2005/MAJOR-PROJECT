"""LangGraph seminar deck: 20 clean, minimal slides (native shapes, editable in Canva)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

HEAD, BODY = "Poppins", "Inter"
INK, MUTED, ACCENT, SOFT, CARD, LINE, WHITE = "1F2937", "6B7280", "4F46E5", "EEF2FF", "F3F4F6", "D1D5DB", "FFFFFF"
FOOT = "LangGraph Seminar  ·  Nikhil Jimmy Thomas (MUT23CA055)"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def tb(s, x, y, w, h, paras, size=18, color=INK, font=BODY, bold=False, align="l", anchor="t", spacing=1.15, after=6):
    """paras: list of str | dict(text|runs, size, bold, color, bullet, font, align, after)."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_paras(tf, paras, size, color, font, bold, align, spacing, after)
    return box


def fill_paras(tf, paras, size, color, font, bold, align, spacing=1.15, after=6):
    for i, p in enumerate(paras):
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p.get("align", align)]
        para.line_spacing = spacing
        para.space_after = Pt(p.get("after", after))
        runs = p.get("runs") or [(p.get("text", ""), {})]
        if p.get("bullet"):
            runs = [("•  ", {"color": ACCENT, "bold": True})] + list(runs)
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", p.get("font", font))
            f.size = Pt(o.get("size", p.get("size", size)))
            f.bold = o.get("bold", p.get("bold", bold))
            f.color.rgb = rgb(o.get("color", p.get("color", color)))


def rect(s, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1.25):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.text_frame.text = ""
    return sh


def node(s, x, y, w, h, label, fill=WHITE, line=ACCENT, size=15, color=INK, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = rect(s, x, y, w, h, fill=fill, line=line, shape=shape, radius=0.18)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_paras(tf, [{"text": t} for t in label.split("\n")], size, color, BODY, bold, "c", 1.0, 0)
    return sh


def arrow(s, x1, y1, x2, y2, color=MUTED, w=1.75, both=False):
    assert x1 == x2 or y1 == y2, "straight connectors only"
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    for tag, on in (("a:headEnd", both), ("a:tailEnd", True)):
        if on:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "med")
            e.set("len", "med")
    return c


def base(title, num, kicker=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0.6, 0.62, 0.09, 0.62, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        tb(s, 0.85, 0.42, 10, 0.3, [kicker.upper()], size=11, color=ACCENT, font=BODY, bold=True)
    tb(s, 0.85, 0.62, 11.5, 0.7, [title], size=30, font=HEAD, bold=True, anchor="m")
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(6.95), Inches(12.73), Inches(6.95)).line.color.rgb = rgb("E5E7EB")
    tb(s, 0.6, 7.02, 9, 0.3, [FOOT], size=10, color=MUTED)
    tb(s, 11.73, 7.02, 1.0, 0.3, [str(num)], size=10, color=MUTED, align="r")
    return s


def bullets(items, size=19):
    return [{"runs": [(h, {"bold": True}), (b, {"color": MUTED})] if b else [(h, {})], "bullet": True, "size": size, "after": 10}
            for h, b in items]


def card(s, x, y, w, h, head, body, num=None, size=14):
    rect(s, x, y, w, h, fill=CARD)
    tx = x + 0.25
    if num is not None:
        c = rect(s, x + 0.25, y + 0.25, 0.5, 0.5, fill=ACCENT, shape=MSO_SHAPE.OVAL)
        fill_paras(c.text_frame, [{"text": str(num), "align": "c"}], 14, WHITE, HEAD, True, "c")
        c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        tx = x + 0.9
    tb(s, tx, y + 0.25, x + w - tx - 0.2, 0.5, [head], size=17, font=HEAD, bold=True, anchor="m")
    tb(s, x + 0.25, y + 0.9, w - 0.5, h - 1.05, body if isinstance(body, list) else [body], size=size, color=MUTED)


def table(s, x, y, colw, rows, rh=0.62, size=13, hsize=14):
    for r, row in enumerate(rows):
        xx = x
        for k, cell in enumerate(row):
            hdr = r == 0
            sh = rect(s, xx, y + r * rh, colw[k] - 0.04, rh - 0.05, fill=ACCENT if hdr else (CARD if r % 2 else WHITE),
                      line=None if hdr else "E5E7EB", shape=MSO_SHAPE.RECTANGLE, lw=0.75)
            tf = sh.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.12)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill_paras(tf, [cell], hsize if hdr else size, WHITE if hdr else INK, BODY, hdr or k == 0, "l", 1.0, 0)
            xx += colw[k]


exec(open("content.py").read())
prs.save("LangGraph-Seminar.pptx")
print("slides:", len(prs.slides))
